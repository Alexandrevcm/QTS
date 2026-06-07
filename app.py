# -*- coding: utf-8 -*-
"""
Quadro Semanal de Instrutores - v1.3

Sistema web simples para:
- cadastrar instrutores;
- cadastrar semanas e horários vagos;
- gerar link público para escolha de horários;
- bloquear vagas preenchidas e conflitos de horário;
- montar o quadro semanal final;
- exportar Excel, PDF e texto para WhatsApp;
- controlar limite semanal de h/a por instrutor;
- permitir inclusão manual de escolhas pelo administrador;
- publicar online com banco compartilhado PostgreSQL/Supabase quando configurado.

Rodar localmente:
    streamlit run app.py
"""

from __future__ import annotations

import io
import sqlite3
import uuid
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

try:
    import psycopg2
    from psycopg2.extras import RealDictCursor

    PSYCOPG2_DISPONIVEL = True
except Exception:
    psycopg2 = None
    RealDictCursor = None
    PSYCOPG2_DISPONIVEL = False

DB_INTEGRITY_ERRORS = (sqlite3.IntegrityError,)
DB_OPERATIONAL_ERRORS = (sqlite3.OperationalError,)
if PSYCOPG2_DISPONIVEL:
    DB_INTEGRITY_ERRORS = DB_INTEGRITY_ERRORS + (psycopg2.IntegrityError,)
    DB_OPERATIONAL_ERRORS = DB_OPERATIONAL_ERRORS + (psycopg2.OperationalError,)

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    REPORTLAB_DISPONIVEL = True
except Exception:
    REPORTLAB_DISPONIVEL = False


APP_NAME = "Quadro Semanal de Instrutores"
APP_VERSION = "1.3"
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "quadro_instrutores.db"
DATA_DIR.mkdir(exist_ok=True)

st.set_page_config(
    page_title=APP_NAME,
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="expanded",
)

DIAS_PT = {
    0: "Segunda-feira",
    1: "Terça-feira",
    2: "Quarta-feira",
    3: "Quinta-feira",
    4: "Sexta-feira",
    5: "Sábado",
    6: "Domingo",
}

STATUS_SEMANA = ["Aberta", "Encerrada", "Publicada"]


# =============================================================================
# Banco de dados
# =============================================================================


def segredo(nome: str, padrao: Any = None) -> Any:
    """Lê secrets do Streamlit Cloud com fallback local."""
    try:
        return st.secrets.get(nome, padrao)
    except Exception:
        return padrao


def obter_database_url() -> str:
    """Aceita várias chaves comuns para facilitar a publicação."""
    for chave in ("DATABASE_URL", "POSTGRES_URL", "SUPABASE_DATABASE_URL"):
        valor = segredo(chave, "")
        if valor:
            return str(valor).strip()

    # Também aceita se o usuário organizar secrets por seção.
    try:
        for secao in ("postgres", "database", "supabase"):
            if secao in st.secrets and st.secrets[secao].get("url"):
                return str(st.secrets[secao]["url"]).strip()
    except Exception:
        pass
    return ""


def db_backend() -> str:
    """
    sqlite  = banco local, bom para testes.
    postgres = banco compartilhado, ideal para Streamlit Cloud/Supabase.
    auto = usa Postgres se houver DATABASE_URL; caso contrário, SQLite.
    """
    modo = str(segredo("DATABASE_MODE", "auto") or "auto").strip().lower()
    url = obter_database_url()
    if modo in {"postgres", "postgresql", "supabase"}:
        return "postgres"
    if modo in {"sqlite", "local"}:
        return "sqlite"
    return "postgres" if url else "sqlite"


def usando_postgres() -> bool:
    return db_backend() == "postgres"


def sql_db(sql: str) -> str:
    """Converte marcadores simples de parâmetro do SQLite para psycopg2."""
    if usando_postgres():
        return sql.replace("?", "%s").replace(" COLLATE NOCASE", "")
    return sql


def conectar():
    if usando_postgres():
        if not PSYCOPG2_DISPONIVEL:
            raise RuntimeError(
                "O modo PostgreSQL foi configurado, mas psycopg2-binary não está instalado. "
                "Confira o requirements.txt."
            )
        url = obter_database_url()
        if not url:
            raise RuntimeError(
                "DATABASE_MODE está como PostgreSQL, mas DATABASE_URL/POSTGRES_URL não foi configurado nos secrets."
            )
        sslmode = str(segredo("POSTGRES_SSLMODE", "require") or "require")
        conn = psycopg2.connect(url, cursor_factory=RealDictCursor, sslmode=sslmode)
        conn.autocommit = False
        return conn

    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA busy_timeout = 5000")
    return conn


@st.cache_resource(show_spinner=False)
def get_conn():
    return conectar()


def executar(sql: str, params: tuple[Any, ...] = ()):  # cursor sqlite ou psycopg2
    conn = get_conn()
    try:
        if usando_postgres():
            cur = conn.cursor()
            cur.execute(sql_db(sql), params)
        else:
            cur = conn.execute(sql, params)
        conn.commit()
        return cur
    except Exception:
        conn.rollback()
        raise


def inserir_retornando_id(sql: str, params: tuple[Any, ...] = ()) -> int:
    """Executa INSERT e retorna o id gerado, tanto no SQLite quanto no Postgres."""
    if usando_postgres():
        sql_limpo = sql.strip().rstrip(";")
        cur = executar(sql_limpo + " RETURNING id", params)
        row = cur.fetchone()
        return int(row["id"] if isinstance(row, dict) else row[0])
    cur = executar(sql, params)
    return int(cur.lastrowid)


def consultar(sql: str, params: tuple[Any, ...] = ()) -> pd.DataFrame:
    conn = get_conn()
    return pd.read_sql_query(sql_db(sql), conn, params=params)


def consultar_um(sql: str, params: tuple[Any, ...] = ()):
    conn = get_conn()
    if usando_postgres():
        cur = conn.cursor()
        cur.execute(sql_db(sql), params)
        return cur.fetchone()
    cur = conn.execute(sql, params)
    return cur.fetchone()


def executar_em_transacao(conn, sql: str, params: tuple[Any, ...] = ()):
    if usando_postgres():
        cur = conn.cursor()
        cur.execute(sql_db(sql), params)
        return cur
    return conn.execute(sql, params)


def iniciar_transacao_imediata(conn) -> None:
    # Antes de abrir uma transação manual, limpamos qualquer transação pendente.
    try:
        conn.rollback()
    except Exception:
        pass
    if usando_postgres():
        cur = conn.cursor()
        cur.execute("BEGIN")
    else:
        conn.execute("BEGIN IMMEDIATE")


def init_db_sqlite() -> None:
    conn = get_conn()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS instrutores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            posto_grad TEXT,
            matricula TEXT UNIQUE,
            telefone TEXT,
            areas TEXT,
            codigo_acesso TEXT,
            ativo INTEGER NOT NULL DEFAULT 1,
            criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS semanas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            token TEXT NOT NULL UNIQUE,
            titulo TEXT NOT NULL,
            data_inicio TEXT NOT NULL,
            data_fim TEXT NOT NULL,
            prazo_escolha TEXT,
            status TEXT NOT NULL DEFAULT 'Aberta',
            observacoes TEXT,
            limite_horas_instrutor INTEGER NOT NULL DEFAULT 10,
            encerrada_em TEXT,
            publicada_em TEXT,
            criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS horarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            semana_id INTEGER NOT NULL,
            data_aula TEXT NOT NULL,
            hora_inicio TEXT NOT NULL,
            hora_fim TEXT NOT NULL,
            disciplina TEXT NOT NULL DEFAULT '',
            local TEXT,
            vagas INTEGER NOT NULL DEFAULT 1,
            habilitacao TEXT,
            carga_horaria INTEGER NOT NULL DEFAULT 0,
            bloqueado INTEGER NOT NULL DEFAULT 0,
            observacoes TEXT,
            criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (semana_id) REFERENCES semanas(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS escolhas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            horario_id INTEGER NOT NULL,
            instrutor_id INTEGER NOT NULL,
            status TEXT NOT NULL DEFAULT 'Confirmada',
            criado_em TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (horario_id) REFERENCES horarios(id) ON DELETE CASCADE,
            FOREIGN KEY (instrutor_id) REFERENCES instrutores(id) ON DELETE CASCADE,
            UNIQUE (horario_id, instrutor_id)
        );

        CREATE INDEX IF NOT EXISTS idx_instrutores_matricula ON instrutores(matricula);
        CREATE INDEX IF NOT EXISTS idx_semanas_token ON semanas(token);
        CREATE INDEX IF NOT EXISTS idx_horarios_semana ON horarios(semana_id);
        CREATE INDEX IF NOT EXISTS idx_escolhas_horario ON escolhas(horario_id);
        CREATE INDEX IF NOT EXISTS idx_escolhas_instrutor ON escolhas(instrutor_id);
        """
    )
    migracoes = [
        "ALTER TABLE horarios ADD COLUMN carga_horaria INTEGER NOT NULL DEFAULT 0",
        "ALTER TABLE instrutores ADD COLUMN codigo_acesso TEXT",
        "ALTER TABLE semanas ADD COLUMN limite_horas_instrutor INTEGER NOT NULL DEFAULT 10",
        "ALTER TABLE semanas ADD COLUMN encerrada_em TEXT",
        "ALTER TABLE semanas ADD COLUMN publicada_em TEXT",
    ]
    for sql in migracoes:
        try:
            conn.execute(sql)
        except sqlite3.OperationalError as exc:
            if "duplicate column name" not in str(exc).lower():
                raise
    conn.commit()


def init_db_postgres() -> None:
    comandos = [
        """
        CREATE TABLE IF NOT EXISTS instrutores (
            id SERIAL PRIMARY KEY,
            nome TEXT NOT NULL,
            posto_grad TEXT,
            matricula TEXT UNIQUE,
            telefone TEXT,
            areas TEXT,
            codigo_acesso TEXT,
            ativo INTEGER NOT NULL DEFAULT 1,
            criado_em TEXT NOT NULL DEFAULT to_char(now(), 'YYYY-MM-DD"T"HH24:MI:SS')
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS semanas (
            id SERIAL PRIMARY KEY,
            token TEXT NOT NULL UNIQUE,
            titulo TEXT NOT NULL,
            data_inicio TEXT NOT NULL,
            data_fim TEXT NOT NULL,
            prazo_escolha TEXT,
            status TEXT NOT NULL DEFAULT 'Aberta',
            observacoes TEXT,
            limite_horas_instrutor INTEGER NOT NULL DEFAULT 10,
            encerrada_em TEXT,
            publicada_em TEXT,
            criado_em TEXT NOT NULL DEFAULT to_char(now(), 'YYYY-MM-DD"T"HH24:MI:SS')
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS horarios (
            id SERIAL PRIMARY KEY,
            semana_id INTEGER NOT NULL REFERENCES semanas(id) ON DELETE CASCADE,
            data_aula TEXT NOT NULL,
            hora_inicio TEXT NOT NULL,
            hora_fim TEXT NOT NULL,
            disciplina TEXT NOT NULL DEFAULT '',
            local TEXT,
            vagas INTEGER NOT NULL DEFAULT 1,
            habilitacao TEXT,
            carga_horaria INTEGER NOT NULL DEFAULT 0,
            bloqueado INTEGER NOT NULL DEFAULT 0,
            observacoes TEXT,
            criado_em TEXT NOT NULL DEFAULT to_char(now(), 'YYYY-MM-DD"T"HH24:MI:SS')
        )
        """,
        """
        CREATE TABLE IF NOT EXISTS escolhas (
            id SERIAL PRIMARY KEY,
            horario_id INTEGER NOT NULL REFERENCES horarios(id) ON DELETE CASCADE,
            instrutor_id INTEGER NOT NULL REFERENCES instrutores(id) ON DELETE CASCADE,
            status TEXT NOT NULL DEFAULT 'Confirmada',
            criado_em TEXT NOT NULL DEFAULT to_char(now(), 'YYYY-MM-DD"T"HH24:MI:SS'),
            UNIQUE (horario_id, instrutor_id)
        )
        """,
        "CREATE INDEX IF NOT EXISTS idx_instrutores_matricula ON instrutores(matricula)",
        "CREATE INDEX IF NOT EXISTS idx_semanas_token ON semanas(token)",
        "CREATE INDEX IF NOT EXISTS idx_horarios_semana ON horarios(semana_id)",
        "CREATE INDEX IF NOT EXISTS idx_escolhas_horario ON escolhas(horario_id)",
        "CREATE INDEX IF NOT EXISTS idx_escolhas_instrutor ON escolhas(instrutor_id)",
        "ALTER TABLE horarios ADD COLUMN IF NOT EXISTS carga_horaria INTEGER NOT NULL DEFAULT 0",
        "ALTER TABLE instrutores ADD COLUMN IF NOT EXISTS codigo_acesso TEXT",
        "ALTER TABLE semanas ADD COLUMN IF NOT EXISTS limite_horas_instrutor INTEGER NOT NULL DEFAULT 10",
        "ALTER TABLE semanas ADD COLUMN IF NOT EXISTS encerrada_em TEXT",
        "ALTER TABLE semanas ADD COLUMN IF NOT EXISTS publicada_em TEXT",
    ]
    conn = get_conn()
    try:
        cur = conn.cursor()
        for comando in comandos:
            cur.execute(comando)
        conn.commit()
    except Exception:
        conn.rollback()
        raise


def init_db() -> None:
    if usando_postgres():
        init_db_postgres()
    else:
        init_db_sqlite()


def status_banco() -> str:
    if usando_postgres():
        return "PostgreSQL/Supabase compartilhado"
    return f"SQLite local ({DB_PATH.name})"

# =============================================================================
# Utilidades
# =============================================================================



def br_date(valor: str | date | None) -> str:
    if not valor:
        return ""
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return datetime.strptime(str(valor)[:10], "%Y-%m-%d").strftime("%d/%m/%Y")

def iso_date(valor: date | datetime | str) -> str:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    return str(valor)[:10]


def hora_str(valor: time | str) -> str:
    if isinstance(valor, time):
        return valor.strftime("%H:%M")
    return str(valor)[:5]


def parse_hora(valor: str) -> time:
    return datetime.strptime(valor[:5], "%H:%M").time()


def formatar_instrutor(row: pd.Series | sqlite3.Row | dict[str, Any]) -> str:
    posto = (row["posto_grad"] or "").strip()
    nome = (row["nome"] or "").strip()
    return f"{posto} {nome}".strip()


def descricao_aula(valor: Any) -> str:
    texto = "" if valor is None else str(valor).strip()
    return texto if texto else "Horário vago"


def periodo_com_carga(hora_inicio: str, hora_fim: str, carga_horaria: Any = 0) -> str:
    periodo = f"{str(hora_inicio)[:5]} às {str(hora_fim)[:5]}"
    try:
        carga = int(carga_horaria or 0)
    except Exception:
        carga = 0
    if carga > 0:
        periodo += f" ({carga}h/a)"
    return periodo


def datas_uteis_da_semana(data_inicio: str | date, data_fim: str | date) -> list[date]:
    inicio = datetime.strptime(str(data_inicio)[:10], "%Y-%m-%d").date() if not isinstance(data_inicio, date) else data_inicio
    fim = datetime.strptime(str(data_fim)[:10], "%Y-%m-%d").date() if not isinstance(data_fim, date) else data_fim
    dias: list[date] = []
    atual = inicio
    while atual <= fim:
        if atual.weekday() <= 4:
            dias.append(atual)
        atual += timedelta(days=1)
    return dias


GRADE_PADRAO_SEMANAL = [
    ("08:00", "09:30", 2),
    ("09:45", "12:00", 3),
    ("14:00", "15:30", 2),
    ("15:45", "18:00", 3),
]


def get_query_param(nome: str, padrao: str = "") -> str:
    try:
        valor = st.query_params.get(nome, padrao)
    except Exception:
        valor = padrao
    if isinstance(valor, list):
        return valor[0] if valor else padrao
    return valor or padrao


def set_query_params(**kwargs: str) -> None:
    try:
        st.query_params.clear()
        for chave, valor in kwargs.items():
            st.query_params[chave] = valor
    except Exception:
        pass


def mostrar_cabecalho(subtitulo: str = "") -> None:
    st.title(APP_NAME)
    if subtitulo:
        st.caption(subtitulo)
    st.caption(f"Versão {APP_VERSION} • Banco: {status_banco()}")


def senha_admin_configurada() -> str:
    # Em produção, coloque ADMIN_PASSWORD no secrets.toml do Streamlit Cloud.
    try:
        return st.secrets.get("ADMIN_PASSWORD", "admin123")
    except Exception:
        return "admin123"


def obter_base_url() -> str:
    padrao = "http://localhost:8501"
    try:
        return st.secrets.get("BASE_URL", padrao).rstrip("/")
    except Exception:
        return padrao


def criar_dados_exemplo() -> None:
    if consultar_um("SELECT id FROM instrutores LIMIT 1"):
        st.warning("Já existem dados cadastrados. Os dados de exemplo não foram duplicados.")
        return

    instrutores = [
        ("João Silva", "Sgt", "1001", "", "APH; Salvamento"),
        ("Carlos Almeida", "Ten", "1002", "", "Incêndio; APH"),
        ("Marcos Pereira", "Cb", "1003", "", "Salvamento"),
        ("André Souza", "Cap", "1004", "", "Coordenação; Incêndio"),
    ]
    for nome, posto, matricula, telefone, areas in instrutores:
        executar(
            """
            INSERT INTO instrutores (nome, posto_grad, matricula, telefone, areas, ativo)
            VALUES (?, ?, ?, ?, ?, 1)
            """,
            (nome, posto, matricula, telefone, areas),
        )

    hoje = date.today()
    inicio = hoje - timedelta(days=hoje.weekday())
    fim = inicio + timedelta(days=4)
    token = uuid.uuid4().hex[:12]
    semana_id = inserir_retornando_id(
        """
        INSERT INTO semanas (token, titulo, data_inicio, data_fim, prazo_escolha, status, observacoes)
        VALUES (?, ?, ?, ?, ?, 'Aberta', ?)
        """,
        (
            token,
            "Semana de Instrução - Exemplo",
            inicio.isoformat(),
            fim.isoformat(),
            datetime.combine(fim, time(18, 0)).isoformat(timespec="minutes"),
            "Dados de exemplo para teste inicial.",
        ),
    )

    aulas = []
    for dia in datas_uteis_da_semana(inicio, fim):
        for h_ini, h_fim, carga in GRADE_PADRAO_SEMANAL:
            aulas.append((dia, h_ini, h_fim, "", "", 1, "", carga))

    for data_aula, h_ini, h_fim, disciplina, local, vagas, hab, carga in aulas:
        executar(
            """
            INSERT INTO horarios
                (semana_id, data_aula, hora_inicio, hora_fim, disciplina, local, vagas, habilitacao, carga_horaria, bloqueado)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
            """,
            (semana_id, data_aula.isoformat(), h_ini, h_fim, disciplina, local, vagas, hab, int(carga)),
        )

    st.success("Dados de exemplo criados com sucesso.")


# =============================================================================
# Consultas principais
# =============================================================================


def listar_instrutores(ativos: bool | None = None) -> pd.DataFrame:
    where = ""
    params: tuple[Any, ...] = ()
    if ativos is True:
        where = "WHERE ativo = 1"
    elif ativos is False:
        where = "WHERE ativo = 0"
    return consultar(
        f"""
        SELECT id, posto_grad, nome, matricula, telefone, areas, codigo_acesso, ativo, criado_em
        FROM instrutores
        {where}
        ORDER BY ativo DESC, nome COLLATE NOCASE
        """,
        params,
    )


def listar_semanas() -> pd.DataFrame:
    return consultar(
        """
        SELECT id, token, titulo, data_inicio, data_fim, prazo_escolha, status, observacoes, limite_horas_instrutor, encerrada_em, publicada_em, criado_em
        FROM semanas
        ORDER BY data_inicio DESC, id DESC
        """
    )


def obter_semana_por_token(token: str) -> sqlite3.Row | None:
    return consultar_um("SELECT * FROM semanas WHERE token = ?", (token,))


def obter_semana(semana_id: int) -> sqlite3.Row | None:
    return consultar_um("SELECT * FROM semanas WHERE id = ?", (semana_id,))


def quadro_semana(semana_id: int) -> pd.DataFrame:
    if usando_postgres():
        agregador = "STRING_AGG(TRIM(COALESCE(i.posto_grad, '') || ' ' || i.nome), '; ' ORDER BY i.nome)"
    else:
        agregador = "GROUP_CONCAT(TRIM(COALESCE(i.posto_grad, '') || ' ' || i.nome), '; ')"
    return consultar(
        f"""
        SELECT
            h.id AS horario_id,
            h.data_aula,
            h.hora_inicio,
            h.hora_fim,
            h.disciplina,
            h.local,
            h.vagas,
            h.habilitacao,
            h.carga_horaria,
            h.bloqueado,
            h.observacoes,
            COALESCE(COUNT(e.id), 0) AS ocupadas,
            {agregador} AS instrutores
        FROM horarios h
        LEFT JOIN escolhas e ON e.horario_id = h.id AND e.status = 'Confirmada'
        LEFT JOIN instrutores i ON i.id = e.instrutor_id
        WHERE h.semana_id = ?
        GROUP BY h.id, h.data_aula, h.hora_inicio, h.hora_fim, h.disciplina, h.local, h.vagas, h.habilitacao, h.carga_horaria, h.bloqueado, h.observacoes
        ORDER BY h.data_aula, h.hora_inicio, h.hora_fim
        """,
        (semana_id,),
    )

def horarios_disponiveis(semana_id: int) -> pd.DataFrame:
    df = quadro_semana(semana_id)
    if df.empty:
        return df
    df["vagas_restantes"] = df["vagas"] - df["ocupadas"]
    return df[(df["bloqueado"] == 0) & (df["vagas_restantes"] > 0)].copy()


def escolhas_do_instrutor(instrutor_id: int, semana_id: int) -> pd.DataFrame:
    return consultar(
        """
        SELECT
            e.id AS escolha_id,
            h.id AS horario_id,
            h.data_aula,
            h.hora_inicio,
            h.hora_fim,
            h.disciplina,
            h.local,
            h.carga_horaria,
            e.criado_em
        FROM escolhas e
        JOIN horarios h ON h.id = e.horario_id
        WHERE e.instrutor_id = ?
          AND h.semana_id = ?
          AND e.status = 'Confirmada'
        ORDER BY h.data_aula, h.hora_inicio
        """,
        (instrutor_id, semana_id),
    )


def obter_instrutor_por_matricula(matricula: str) -> sqlite3.Row | None:
    return consultar_um("SELECT * FROM instrutores WHERE matricula = ?", (matricula.strip(),))


def obter_instrutor(instrutor_id: int) -> sqlite3.Row | None:
    return consultar_um("SELECT * FROM instrutores WHERE id = ?", (instrutor_id,))


def contar_ocupadas(horario_id: int) -> int:
    row = consultar_um(
        "SELECT COUNT(*) AS total FROM escolhas WHERE horario_id = ? AND status = 'Confirmada'",
        (horario_id,),
    )
    return int(row["total"] if row else 0)


def carga_instrutor_semana(instrutor_id: int, semana_id: int) -> int:
    row = consultar_um(
        """
        SELECT COALESCE(SUM(COALESCE(h.carga_horaria, 0)), 0) AS total
        FROM escolhas e
        JOIN horarios h ON h.id = e.horario_id
        WHERE e.instrutor_id = ?
          AND h.semana_id = ?
          AND e.status = 'Confirmada'
        """,
        (instrutor_id, semana_id),
    )
    return int(row["total"] if row and row["total"] is not None else 0)


def carga_por_instrutor_semana(semana_id: int) -> pd.DataFrame:
    return consultar(
        """
        SELECT
            i.id AS instrutor_id,
            TRIM(COALESCE(i.posto_grad, '') || ' ' || i.nome) AS instrutor,
            i.matricula,
            COUNT(e.id) AS quantidade_horarios,
            COALESCE(SUM(COALESCE(h.carga_horaria, 0)), 0) AS carga_total
        FROM escolhas e
        JOIN horarios h ON h.id = e.horario_id
        JOIN instrutores i ON i.id = e.instrutor_id
        WHERE h.semana_id = ?
          AND e.status = 'Confirmada'
        GROUP BY i.id
        ORDER BY carga_total DESC, instrutor COLLATE NOCASE
        """,
        (semana_id,),
    )


def limite_horas_da_semana(semana: sqlite3.Row | dict[str, Any] | pd.Series) -> int:
    try:
        return int(semana["limite_horas_instrutor"] or 0)
    except Exception:
        return 0


def ha_conflito_horario(instrutor_id: int, horario_id: int) -> tuple[bool, str]:
    horario = consultar_um("SELECT * FROM horarios WHERE id = ?", (horario_id,))
    if not horario:
        return True, "Horário não encontrado."

    existentes = consultar(
        """
        SELECT h.data_aula, h.hora_inicio, h.hora_fim, h.disciplina
        FROM escolhas e
        JOIN horarios h ON h.id = e.horario_id
        WHERE e.instrutor_id = ?
          AND h.semana_id = ?
          AND h.data_aula = ?
          AND e.status = 'Confirmada'
        """,
        (instrutor_id, horario["semana_id"], horario["data_aula"]),
    )

    novo_ini = parse_hora(horario["hora_inicio"])
    novo_fim = parse_hora(horario["hora_fim"])
    for _, item in existentes.iterrows():
        ini = parse_hora(item["hora_inicio"])
        fim = parse_hora(item["hora_fim"])
        sobrepoe = novo_ini < fim and novo_fim > ini
        if sobrepoe:
            desc = f'{descricao_aula(item["disciplina"])} ({item["hora_inicio"]} às {item["hora_fim"]})'
            return True, f"Conflito com outro horário já escolhido: {desc}."
    return False, ""


def escolher_horario(instrutor_id: int, horario_id: int, origem_admin: bool = False) -> tuple[bool, str]:
    conn = get_conn()
    try:
        iniciar_transacao_imediata(conn)

        horario = executar_em_transacao(conn, "SELECT * FROM horarios WHERE id = ?", (horario_id,)).fetchone()
        if not horario:
            conn.rollback()
            return False, "Horário não encontrado."
        if int(horario["bloqueado"]) == 1:
            conn.rollback()
            return False, "Este horário está bloqueado pelo administrador."

        semana = executar_em_transacao(conn, "SELECT * FROM semanas WHERE id = ?", (horario["semana_id"],)).fetchone()
        if not semana:
            conn.rollback()
            return False, "Semana não encontrada."
        if semana["status"] != "Aberta" and not origem_admin:
            conn.rollback()
            return False, "A semana não está aberta para escolhas."

        if semana["prazo_escolha"] and not origem_admin:
            prazo = datetime.fromisoformat(str(semana["prazo_escolha"]))
            if datetime.now() > prazo:
                conn.rollback()
                return False, "O prazo de escolha já terminou."

        ocupadas_row = executar_em_transacao(
            conn,
            "SELECT COUNT(*) AS total FROM escolhas WHERE horario_id = ? AND status = 'Confirmada'",
            (horario_id,),
        ).fetchone()
        ocupadas = ocupadas_row["total"] if ocupadas_row else 0
        if int(ocupadas) >= int(horario["vagas"]):
            conn.rollback()
            return False, "As vagas desse horário acabaram."

        ja_escolheu = executar_em_transacao(
            conn,
            """
            SELECT id FROM escolhas
            WHERE horario_id = ? AND instrutor_id = ? AND status = 'Confirmada'
            """,
            (horario_id, instrutor_id),
        ).fetchone()
        if ja_escolheu:
            conn.rollback()
            return False, "Você já escolheu esse horário."

        limite = int(semana["limite_horas_instrutor"] or 0)
        carga_nova = int(horario["carga_horaria"] or 0)
        if limite > 0 and carga_nova > 0:
            carga_row = executar_em_transacao(
                conn,
                """
                SELECT COALESCE(SUM(COALESCE(h.carga_horaria, 0)), 0) AS total
                FROM escolhas e
                JOIN horarios h ON h.id = e.horario_id
                WHERE e.instrutor_id = ?
                  AND h.semana_id = ?
                  AND e.status = 'Confirmada'
                """,
                (instrutor_id, horario["semana_id"]),
            ).fetchone()
            carga_atual = carga_row["total"] if carga_row else 0
            if int(carga_atual or 0) + carga_nova > limite:
                conn.rollback()
                return (
                    False,
                    f"Limite semanal ultrapassado. Você já possui {int(carga_atual or 0)}h/a; "
                    f"este horário tem {carga_nova}h/a; limite da semana: {limite}h/a."
                )

        existentes = executar_em_transacao(
            conn,
            """
            SELECT h.data_aula, h.hora_inicio, h.hora_fim, h.disciplina
            FROM escolhas e
            JOIN horarios h ON h.id = e.horario_id
            WHERE e.instrutor_id = ?
              AND h.semana_id = ?
              AND h.data_aula = ?
              AND e.status = 'Confirmada'
            """,
            (instrutor_id, horario["semana_id"], horario["data_aula"]),
        ).fetchall()
        novo_ini = parse_hora(horario["hora_inicio"])
        novo_fim = parse_hora(horario["hora_fim"])
        for item in existentes:
            ini = parse_hora(item["hora_inicio"])
            fim = parse_hora(item["hora_fim"])
            if novo_ini < fim and novo_fim > ini:
                conn.rollback()
                return False, f"Conflito com outro horário já escolhido: {descricao_aula(item['disciplina'])} ({item['hora_inicio']} às {item['hora_fim']})."

        executar_em_transacao(
            conn,
            """
            INSERT INTO escolhas (horario_id, instrutor_id, status)
            VALUES (?, ?, 'Confirmada')
            """,
            (horario_id, instrutor_id),
        )
        conn.commit()
        return True, "Horário escolhido com sucesso."
    except DB_INTEGRITY_ERRORS as exc:
        conn.rollback()
        return False, f"Não foi possível registrar a escolha: {exc}."
    except Exception as exc:
        conn.rollback()
        return False, f"Erro ao registrar escolha: {exc}."


# =============================================================================
# Exportações
# =============================================================================



def preparar_df_quadro(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    saida = df.copy()
    saida["Data"] = pd.to_datetime(saida["data_aula"]).dt.strftime("%d/%m/%Y")
    saida["Dia"] = pd.to_datetime(saida["data_aula"]).dt.weekday.map(DIAS_PT)
    saida["Horário"] = saida.apply(lambda r: periodo_com_carga(r["hora_inicio"], r["hora_fim"], r.get("carga_horaria", 0)), axis=1)
    saida["Descrição"] = saida["disciplina"].apply(descricao_aula)
    saida["Vagas"] = saida["ocupadas"].astype(str) + "/" + saida["vagas"].astype(str)
    saida["Instrutores"] = saida["instrutores"].fillna("—")
    saida["Bloqueado"] = saida["bloqueado"].map({0: "Não", 1: "Sim"})
    return saida[
        [
            "Dia",
            "Data",
            "Horário",
            "Descrição",
            "local",
            "Instrutores",
            "Vagas",
            "habilitacao",
            "Bloqueado",
            "observacoes",
        ]
    ].rename(
        columns={
            "local": "Local",
            "habilitacao": "Habilitação exigida",
            "observacoes": "Observações",
        }
    )

def exportar_excel(df: pd.DataFrame, titulo: str) -> bytes:
    """
    Exporta o quadro para Excel.

    Correção v1.2.1:
    O cabeçalho do título usa células mescladas. Em células mescladas, o openpyxl
    pode devolver objetos MergedCell, que não possuem o atributo column_letter.
    Por isso, a largura das colunas passa a ser calculada pelo índice da coluna
    com get_column_letter(), e não por col[0].column_letter.
    """
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Quadro")
        ws = writer.book["Quadro"]
        ws.insert_rows(1)
        ws["A1"] = titulo
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(df.columns)))

        for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
            max_len = 0
            for cell in col:
                valor = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(valor))
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = min(max_len + 2, 45)

    buffer.seek(0)
    return buffer.getvalue()


def exportar_pdf(df: pd.DataFrame, titulo: str) -> bytes:
    if not REPORTLAB_DISPONIVEL:
        raise RuntimeError("Biblioteca reportlab não está instalada.")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    elementos = [Paragraph(titulo, styles["Title"]), Spacer(1, 12)]

    tabela_dados = [list(df.columns)] + df.fillna("—").astype(str).values.tolist()
    tabela = Table(tabela_dados, repeatRows=1)
    tabela.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    elementos.append(tabela)
    doc.build(elementos)
    buffer.seek(0)
    return buffer.getvalue()



def gerar_texto_whatsapp(df: pd.DataFrame, semana: sqlite3.Row) -> str:
    if df.empty:
        return "Nenhum horário cadastrado."

    linhas = [
        f"*{semana['titulo']}*",
        f"Período: {br_date(semana['data_inicio'])} a {br_date(semana['data_fim'])}",
        "",
    ]
    df2 = df.sort_values(["data_aula", "hora_inicio"])
    data_atual = None
    for _, row in df2.iterrows():
        data = row["data_aula"]
        if data != data_atual:
            data_atual = data
            dt = datetime.strptime(data, "%Y-%m-%d").date()
            linhas.append(f"*{DIAS_PT[dt.weekday()]} - {br_date(dt)}*")
        instrutores = row["instrutores"] if pd.notna(row["instrutores"]) and row["instrutores"] else "Pendente"
        local = f" - {row['local']}" if row["local"] else ""
        descricao = descricao_aula(row.get("disciplina", ""))
        periodo = periodo_com_carga(row["hora_inicio"], row["hora_fim"], row.get("carga_horaria", 0))
        linhas.append(f"{periodo} | {descricao}{local} | {instrutores}")
    return "\n".join(linhas)

# =============================================================================
# Componentes de interface
# =============================================================================


def selecionar_semana(label: str = "Semana") -> int | None:
    semanas = listar_semanas()
    if semanas.empty:
        st.info("Nenhuma semana cadastrada ainda.")
        return None

    opcoes = {}
    for _, row in semanas.iterrows():
        nome = f"{row['titulo']} — {br_date(row['data_inicio'])} a {br_date(row['data_fim'])} — {row['status']}"
        opcoes[nome] = int(row["id"])

    escolha = st.selectbox(label, list(opcoes.keys()))
    return opcoes[escolha]


def formulario_login_admin() -> bool:
    if st.session_state.get("admin_logado"):
        return True

    st.subheader("Acesso administrativo")
    st.info("Senha padrão inicial: admin123. Altere no arquivo `.streamlit/secrets.toml` antes de publicar o app.")
    with st.form("login_admin"):
        senha = st.text_input("Senha", type="password")
        entrar = st.form_submit_button("Entrar")
    if entrar:
        if senha == senha_admin_configurada():
            st.session_state["admin_logado"] = True
            st.success("Acesso liberado.")
            st.rerun()
        else:
            st.error("Senha incorreta.")
    return False


# =============================================================================
# Páginas públicas
# =============================================================================



def pagina_escolha_publica(token: str) -> None:
    mostrar_cabecalho("Link público para escolha de horários pelos instrutores")

    semana = obter_semana_por_token(token)
    if not semana:
        st.error("Link inválido ou semana não encontrada.")
        return

    st.subheader(semana["titulo"])
    st.write(f"**Período:** {br_date(semana['data_inicio'])} a {br_date(semana['data_fim'])}")
    st.write(f"**Status:** {semana['status']}")
    if semana["prazo_escolha"]:
        st.write(f"**Prazo para escolha:** {datetime.fromisoformat(semana['prazo_escolha']).strftime('%d/%m/%Y %H:%M')}")
    if semana["observacoes"]:
        st.info(semana["observacoes"])

    if semana["status"] != "Aberta":
        st.warning("Esta semana não está aberta para novas escolhas.")
        st.stop()

    if semana["prazo_escolha"] and datetime.now() > datetime.fromisoformat(semana["prazo_escolha"]):
        st.warning("O prazo para escolha terminou.")
        st.stop()

    st.divider()
    st.subheader("Identificação do instrutor")

    with st.form("identificacao_instrutor"):
        col1, col2, col3 = st.columns([1, 2, 1])
        with col1:
            posto_grad = st.text_input("Posto/Graduação", placeholder="Ex.: Sgt")
        with col2:
            nome = st.text_input("Nome de guerra ou nome completo *")
        with col3:
            matricula = st.text_input("Matrícula/Identificação *")
        telefone = st.text_input("Telefone/WhatsApp", placeholder="Opcional")
        codigo_acesso = st.text_input(
            "Código de acesso",
            type="password",
            placeholder="Se foi informado pelo administrador",
            help="Opcional. Se o instrutor já tiver código cadastrado, será necessário informá-lo para acessar.",
        )
        confirmar = st.form_submit_button("Continuar")

    instrutor_id = st.session_state.get(f"instrutor_id_{token}")

    if confirmar:
        nome = nome.strip()
        matricula = matricula.strip()
        if not nome or not matricula:
            st.error("Informe pelo menos o nome e a matrícula/identificação.")
            return

        existente = obter_instrutor_por_matricula(matricula)
        codigo_digitado = codigo_acesso.strip()
        if existente:
            codigo_cadastrado = (existente["codigo_acesso"] or "").strip()
            if codigo_cadastrado and codigo_digitado != codigo_cadastrado:
                st.error("Código de acesso incorreto para esta matrícula.")
                return
            instrutor_id = int(existente["id"])
            executar(
                """
                UPDATE instrutores
                SET nome = ?, posto_grad = ?, telefone = ?, ativo = 1
                WHERE id = ?
                """,
                (nome, posto_grad.strip(), telefone.strip(), instrutor_id),
            )
        else:
            instrutor_id = inserir_retornando_id(
                """
                INSERT INTO instrutores (nome, posto_grad, matricula, telefone, codigo_acesso, ativo)
                VALUES (?, ?, ?, ?, ?, 1)
                """,
                (nome, posto_grad.strip(), matricula, telefone.strip(), codigo_digitado or None),
            )

        st.session_state[f"instrutor_id_{token}"] = instrutor_id
        st.success("Identificação registrada. Agora escolha seus horários.")
        st.rerun()

    if not instrutor_id:
        st.stop()

    instrutor = obter_instrutor(int(instrutor_id))
    if not instrutor:
        st.error("Instrutor não encontrado. Preencha a identificação novamente.")
        return

    st.success(f"Instrutor identificado: {formatar_instrutor(instrutor)}")
    limite_semana = limite_horas_da_semana(semana)
    carga_atual_publica = carga_instrutor_semana(int(instrutor_id), int(semana["id"]))
    col_carga1, col_carga2 = st.columns(2)
    col_carga1.metric("Minha carga escolhida", f"{carga_atual_publica}h/a")
    col_carga2.metric("Limite da semana", "Sem limite" if limite_semana <= 0 else f"{limite_semana}h/a")

    st.divider()
    st.subheader("Meus horários escolhidos")
    minhas = escolhas_do_instrutor(int(instrutor_id), int(semana["id"]))
    if minhas.empty:
        st.caption("Você ainda não escolheu horários nesta semana.")
    else:
        minhas_view = minhas.assign(
            Data=lambda d: pd.to_datetime(d["data_aula"]).dt.strftime("%d/%m/%Y"),
            Horário=lambda d: d.apply(lambda r: periodo_com_carga(r["hora_inicio"], r["hora_fim"], r.get("carga_horaria", 0)), axis=1),
            Descrição=lambda d: d["disciplina"].apply(descricao_aula),
        )
        st.dataframe(
            minhas_view[["Data", "Horário", "Descrição", "local"]].rename(columns={"local": "Local"}),
            use_container_width=True,
            hide_index=True,
        )
        escolha_cancelar = st.selectbox(
            "Cancelar uma escolha",
            ["Não cancelar"]
            + [
                f"{row['escolha_id']} | {br_date(row['data_aula'])} {periodo_com_carga(row['hora_inicio'], row['hora_fim'], row.get('carga_horaria', 0))} - {descricao_aula(row['disciplina'])}"
                for _, row in minhas.iterrows()
            ],
        )
        if escolha_cancelar != "Não cancelar" and st.button("Cancelar escolha selecionada"):
            escolha_id = int(escolha_cancelar.split("|")[0].strip())
            executar("DELETE FROM escolhas WHERE id = ? AND instrutor_id = ?", (escolha_id, int(instrutor_id)))
            st.success("Escolha cancelada.")
            st.rerun()

    st.divider()
    st.subheader("Horários disponíveis")
    disponiveis = horarios_disponiveis(int(semana["id"]))
    if disponiveis.empty:
        st.warning("Não há horários disponíveis no momento.")
        return

    for _, row in disponiveis.iterrows():
        with st.container(border=True):
            col1, col2 = st.columns([4, 1])
            with col1:
                data_aula = datetime.strptime(row["data_aula"], "%Y-%m-%d").date()
                periodo = periodo_com_carga(row["hora_inicio"], row["hora_fim"], row.get("carga_horaria", 0))
                st.markdown(f"**{DIAS_PT[data_aula.weekday()]} - {br_date(data_aula)} | {periodo}**")
                st.write(f"**Descrição:** {descricao_aula(row['disciplina'])}")
                detalhes = []
                if row["local"]:
                    detalhes.append(f"Local: {row['local']}")
                if row["habilitacao"]:
                    detalhes.append(f"Habilitação: {row['habilitacao']}")
                detalhes.append(f"Vagas restantes: {int(row['vagas_restantes'])}")
                st.caption(" | ".join(detalhes))
                if row["observacoes"]:
                    st.caption(f"Observações: {row['observacoes']}")
            with col2:
                carga_nova = int(row.get("carga_horaria", 0) or 0)
                excede_limite = limite_semana > 0 and carga_nova > 0 and (carga_atual_publica + carga_nova > limite_semana)
                if excede_limite:
                    st.caption("Ultrapassa seu limite semanal")
                if st.button("Escolher", key=f"escolher_{row['horario_id']}", disabled=excede_limite):
                    ok, msg = escolher_horario(int(instrutor_id), int(row["horario_id"]))
                    if ok:
                        st.success(msg)
                    else:
                        st.error(msg)
                    st.rerun()

# =============================================================================
# Páginas administrativas
# =============================================================================


def pagina_admin() -> None:
    mostrar_cabecalho("Administração do quadro semanal")
    if not formulario_login_admin():
        return

    with st.sidebar:
        st.success("Administrador logado")
        if st.button("Sair"):
            st.session_state["admin_logado"] = False
            st.rerun()
        st.divider()
        if st.button("Criar dados de exemplo"):
            criar_dados_exemplo()
            st.rerun()

    abas = st.tabs(
        [
            "1. Semanas",
            "2. Instrutores",
            "3. Horários vagos",
            "4. Escolhas",
            "5. Quadro final",
            "6. Link público",
            "7. Ambiente/Nuvem",
        ]
    )

    with abas[0]:
        aba_semanas()
    with abas[1]:
        aba_instrutores()
    with abas[2]:
        aba_horarios()
    with abas[3]:
        aba_escolhas()
    with abas[4]:
        aba_quadro_final()
    with abas[5]:
        aba_link_publico()
    with abas[6]:
        aba_ambiente_nuvem()



def aba_ambiente_nuvem() -> None:
    st.subheader("Ambiente e publicação online")
    st.write("Esta tela ajuda a conferir se o app está pronto para rodar por link público com banco compartilhado.")

    col1, col2 = st.columns(2)
    col1.metric("Modo de banco", status_banco())
    col2.metric("Versão", APP_VERSION)

    st.write("**Endereço base configurado:**")
    st.code(obter_base_url(), language="text")

    if usando_postgres():
        st.success("O app está configurado para usar PostgreSQL/Supabase. As escolhas feitas pelos instrutores serão salvas no banco compartilhado.")
    else:
        st.warning(
            "O app está usando SQLite local. Isso é ótimo para teste, mas não serve como banco compartilhado para instrutores acessando pela internet. "
            "Para publicar de verdade, configure DATABASE_MODE='postgres' e DATABASE_URL nos Secrets do Streamlit Cloud."
        )

    st.divider()
    st.markdown("""
**Secrets mínimos para publicação no Streamlit Cloud:**

```toml
ADMIN_PASSWORD = "troque-esta-senha"
BASE_URL = "https://seu-app.streamlit.app"
DATABASE_MODE = "postgres"
DATABASE_URL = "postgresql://postgres.xxxxx:SUA_SENHA@aws-0-sa-east-1.pooler.supabase.com:6543/postgres"
POSTGRES_SSLMODE = "require"
```

Depois de configurar, reinicie/redeploy o app no Streamlit Cloud.
""")

    try:
        qtd_instrutores = consultar_um("SELECT COUNT(*) AS total FROM instrutores")
        qtd_semanas = consultar_um("SELECT COUNT(*) AS total FROM semanas")
        qtd_horarios = consultar_um("SELECT COUNT(*) AS total FROM horarios")
        qtd_escolhas = consultar_um("SELECT COUNT(*) AS total FROM escolhas")
        st.write("**Resumo do banco atual:**")
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Instrutores", int(qtd_instrutores["total"]))
        c2.metric("Semanas", int(qtd_semanas["total"]))
        c3.metric("Horários", int(qtd_horarios["total"]))
        c4.metric("Escolhas", int(qtd_escolhas["total"]))
    except Exception as exc:
        st.error(f"Não foi possível consultar o banco atual: {exc}")


def aba_semanas() -> None:
    st.subheader("Cadastro de semanas")

    with st.form("nova_semana"):
        col1, col2 = st.columns(2)
        with col1:
            titulo = st.text_input("Título da semana", placeholder="Ex.: Semana de Instrução - 10 a 14/06")
            data_inicio = st.date_input("Data inicial", value=date.today())
        with col2:
            data_fim = st.date_input("Data final", value=date.today() + timedelta(days=4))
            prazo_data = st.date_input("Data limite para escolha", value=date.today() + timedelta(days=3))
            prazo_hora = st.time_input("Hora limite", value=time(18, 0))
            limite_horas = st.number_input("Limite h/a por instrutor na semana", min_value=0, max_value=80, value=10, step=1, help="Use 0 para não limitar.")
        observacoes = st.text_area("Observações", placeholder="Orientações que aparecerão no link dos instrutores")
        salvar = st.form_submit_button("Criar semana")

    if salvar:
        if not titulo.strip():
            st.error("Informe um título.")
        elif data_fim < data_inicio:
            st.error("A data final não pode ser anterior à data inicial.")
        else:
            token = uuid.uuid4().hex[:12]
            prazo = datetime.combine(prazo_data, prazo_hora).isoformat(timespec="minutes")
            executar(
                """
                INSERT INTO semanas (token, titulo, data_inicio, data_fim, prazo_escolha, status, observacoes, limite_horas_instrutor)
                VALUES (?, ?, ?, ?, ?, 'Aberta', ?, ?)
                """,
                (token, titulo.strip(), iso_date(data_inicio), iso_date(data_fim), prazo, observacoes.strip(), int(limite_horas)),
            )
            st.success("Semana criada.")
            st.rerun()

    st.divider()
    semanas = listar_semanas()
    if semanas.empty:
        st.info("Nenhuma semana cadastrada.")
        return

    st.dataframe(
        semanas.assign(
            Início=lambda d: pd.to_datetime(d["data_inicio"]).dt.strftime("%d/%m/%Y"),
            Fim=lambda d: pd.to_datetime(d["data_fim"]).dt.strftime("%d/%m/%Y"),
            Limite=lambda d: d["limite_horas_instrutor"].fillna(0).astype(int).astype(str) + "h/a",
        )[["id", "titulo", "Início", "Fim", "status", "Limite", "token", "observacoes"]].rename(
            columns={"id": "ID", "titulo": "Título", "status": "Status", "token": "Token", "observacoes": "Observações"}
        ),
        use_container_width=True,
        hide_index=True,
    )

    st.subheader("Editar status da semana")
    semana_id = selecionar_semana("Selecione a semana para editar")
    if semana_id:
        semana = obter_semana(semana_id)
        if semana:
            col1, col2, col3 = st.columns([1, 2, 1])
            with col1:
                novo_status = st.selectbox("Status", STATUS_SEMANA, index=STATUS_SEMANA.index(semana["status"]))
            with col2:
                novo_titulo = st.text_input("Título", value=semana["titulo"])
            with col3:
                novo_limite = st.number_input("Limite h/a por instrutor", min_value=0, max_value=80, value=limite_horas_da_semana(semana), step=1)
            nova_obs = st.text_area("Observações", value=semana["observacoes"] or "")
            col_salvar, col_encerrar, col_publicar, col_excluir = st.columns(4)
            with col_salvar:
                salvar_semana = st.button("Salvar alterações da semana")
            with col_encerrar:
                encerrar_semana = st.button("Encerrar semana")
            with col_publicar:
                publicar_semana = st.button("Publicar semana")
            with col_excluir:
                excluir_semana = st.button("Excluir semana", type="secondary")

            if salvar_semana:
                agora = datetime.now().isoformat(timespec="minutes")
                encerrada_em = semana["encerrada_em"]
                publicada_em = semana["publicada_em"]
                if novo_status == "Encerrada" and not encerrada_em:
                    encerrada_em = agora
                if novo_status == "Publicada" and not publicada_em:
                    publicada_em = agora
                executar(
                    """
                    UPDATE semanas
                    SET status = ?, titulo = ?, observacoes = ?, limite_horas_instrutor = ?, encerrada_em = ?, publicada_em = ?
                    WHERE id = ?
                    """,
                    (novo_status, novo_titulo.strip(), nova_obs.strip(), int(novo_limite), encerrada_em, publicada_em, semana_id),
                )
                st.success("Semana atualizada.")
                st.rerun()
            if encerrar_semana:
                executar("UPDATE semanas SET status = 'Encerrada', encerrada_em = ? WHERE id = ?", (datetime.now().isoformat(timespec="minutes"), semana_id))
                st.success("Semana encerrada. O link público deixa de aceitar novas escolhas.")
                st.rerun()
            if publicar_semana:
                executar("UPDATE semanas SET status = 'Publicada', publicada_em = ? WHERE id = ?", (datetime.now().isoformat(timespec="minutes"), semana_id))
                st.success("Semana publicada.")
                st.rerun()
            if excluir_semana:
                executar("DELETE FROM semanas WHERE id = ?", (semana_id,))
                st.warning("Semana excluída.")
                st.rerun()


def aba_instrutores() -> None:
    st.subheader("Cadastro de instrutores")

    with st.form("novo_instrutor"):
        col1, col2, col3 = st.columns([1, 2, 1])
        with col1:
            posto = st.text_input("Posto/Graduação")
        with col2:
            nome = st.text_input("Nome *")
        with col3:
            matricula = st.text_input("Matrícula/Identificação")
        telefone = st.text_input("Telefone/WhatsApp")
        areas = st.text_input("Áreas/Habilitações", placeholder="Ex.: APH; Incêndio; Salvamento")
        codigo_acesso = st.text_input("Código de acesso", placeholder="Opcional. Ex.: últimos 4 dígitos da matrícula")
        salvar = st.form_submit_button("Cadastrar instrutor")

    if salvar:
        if not nome.strip():
            st.error("Informe o nome.")
        else:
            try:
                executar(
                    """
                    INSERT INTO instrutores (nome, posto_grad, matricula, telefone, areas, codigo_acesso, ativo)
                    VALUES (?, ?, ?, ?, ?, ?, 1)
                    """,
                    (nome.strip(), posto.strip(), matricula.strip() or None, telefone.strip(), areas.strip(), codigo_acesso.strip() or None),
                )
                st.success("Instrutor cadastrado.")
                st.rerun()
            except DB_INTEGRITY_ERRORS:
                st.error("Já existe instrutor com essa matrícula/identificação.")

    instrutores = listar_instrutores()
    if instrutores.empty:
        st.info("Nenhum instrutor cadastrado.")
        return

    st.dataframe(
        instrutores.rename(
            columns={
                "id": "ID",
                "posto_grad": "Posto/Grad",
                "nome": "Nome",
                "matricula": "Matrícula",
                "telefone": "Telefone",
                "areas": "Áreas",
                "codigo_acesso": "Código acesso",
                "ativo": "Ativo",
            }
        ),
        use_container_width=True,
        hide_index=True,
    )

    st.subheader("Editar instrutor")
    opcoes = {f"{row['id']} - {formatar_instrutor(row)}": int(row["id"]) for _, row in instrutores.iterrows()}
    escolha = st.selectbox("Instrutor", list(opcoes.keys()))
    instrutor_id = opcoes[escolha]
    row = obter_instrutor(instrutor_id)
    if row:
        with st.form("editar_instrutor"):
            col1, col2, col3 = st.columns([1, 2, 1])
            with col1:
                novo_posto = st.text_input("Posto/Graduação", value=row["posto_grad"] or "")
            with col2:
                novo_nome = st.text_input("Nome", value=row["nome"] or "")
            with col3:
                nova_matricula = st.text_input("Matrícula/Identificação", value=row["matricula"] or "")
            novo_telefone = st.text_input("Telefone/WhatsApp", value=row["telefone"] or "")
            novas_areas = st.text_input("Áreas/Habilitações", value=row["areas"] or "")
            novo_codigo = st.text_input("Código de acesso", value=row["codigo_acesso"] or "", help="Se preenchido, o instrutor precisará informar este código no link público.")
            novo_ativo = st.checkbox("Ativo", value=bool(row["ativo"]))
            salvar_instrutor = st.form_submit_button("Salvar alterações do instrutor")
        if salvar_instrutor:
            if not novo_nome.strip():
                st.error("Informe o nome do instrutor.")
            else:
                try:
                    executar(
                        """
                        UPDATE instrutores
                        SET posto_grad = ?, nome = ?, matricula = ?, telefone = ?, areas = ?, codigo_acesso = ?, ativo = ?
                        WHERE id = ?
                        """,
                        (
                            novo_posto.strip(),
                            novo_nome.strip(),
                            nova_matricula.strip() or None,
                            novo_telefone.strip(),
                            novas_areas.strip(),
                            novo_codigo.strip() or None,
                            1 if novo_ativo else 0,
                            instrutor_id,
                        ),
                    )
                    st.success("Instrutor atualizado.")
                    st.rerun()
                except DB_INTEGRITY_ERRORS:
                    st.error("Já existe outro instrutor com essa matrícula/identificação.")



def aba_horarios() -> None:
    st.subheader("Cadastro de horários vagos")
    semana_id = selecionar_semana("Semana para cadastrar horários")
    if not semana_id:
        return

    semana = obter_semana(semana_id)
    assert semana is not None

    st.caption(f"Semana selecionada: {semana['titulo']} — {br_date(semana['data_inicio'])} a {br_date(semana['data_fim'])}")
    st.info(
        "Nesta versão, o horário nasce vago. O instrutor escolhe o período disponível. "
        "A descrição/atividade é opcional e pode ser preenchida depois pelo administrador."
    )

    st.markdown("### Gerar grade padrão da semana")
    st.caption(
        "Cria automaticamente os 4 períodos de segunda a sexta: "
        "08:00-09:30 (2h/a), 09:45-12:00 (3h/a), 14:00-15:30 (2h/a) e 15:45-18:00 (3h/a)."
    )
    with st.form("gerar_grade_padrao"):
        col1, col2, col3 = st.columns(3)
        with col1:
            vagas_padrao = st.number_input("Vagas por horário", min_value=1, max_value=20, value=1, step=1)
        with col2:
            local_padrao = st.text_input("Local padrão", placeholder="Opcional")
        with col3:
            bloquear_criados = st.checkbox("Criar bloqueados", value=False)
        descricao_padrao = st.text_input("Descrição padrão", placeholder="Opcional. Deixe em branco para horário vago.")
        gerar = st.form_submit_button("Gerar grade padrão de segunda a sexta")

    if gerar:
        criados = 0
        ignorados = 0
        for data_aula in datas_uteis_da_semana(semana["data_inicio"], semana["data_fim"]):
            for h_ini, h_fim, carga in GRADE_PADRAO_SEMANAL:
                existe = consultar_um(
                    """
                    SELECT id FROM horarios
                    WHERE semana_id = ? AND data_aula = ? AND hora_inicio = ? AND hora_fim = ?
                    """,
                    (semana_id, data_aula.isoformat(), h_ini, h_fim),
                )
                if existe:
                    ignorados += 1
                    continue
                executar(
                    """
                    INSERT INTO horarios
                        (semana_id, data_aula, hora_inicio, hora_fim, disciplina, local, vagas, habilitacao, carga_horaria, bloqueado, observacoes)
                    VALUES (?, ?, ?, ?, ?, ?, ?, '', ?, ?, '')
                    """,
                    (
                        semana_id,
                        data_aula.isoformat(),
                        h_ini,
                        h_fim,
                        descricao_padrao.strip(),
                        local_padrao.strip(),
                        int(vagas_padrao),
                        int(carga),
                        1 if bloquear_criados else 0,
                    ),
                )
                criados += 1
        st.success(f"Grade gerada: {criados} horários criados. {ignorados} horários já existiam e foram ignorados.")
        st.rerun()

    st.divider()
    st.markdown("### Adicionar horário individual")
    with st.form("novo_horario"):
        col1, col2, col3, col4, col5 = st.columns(5)
        with col1:
            data_aula = st.date_input(
                "Data",
                value=datetime.strptime(semana["data_inicio"], "%Y-%m-%d").date(),
                min_value=datetime.strptime(semana["data_inicio"], "%Y-%m-%d").date(),
                max_value=datetime.strptime(semana["data_fim"], "%Y-%m-%d").date(),
            )
        with col2:
            hora_inicio = st.time_input("Início", value=time(8, 0))
        with col3:
            hora_fim = st.time_input("Fim", value=time(9, 30))
        with col4:
            carga_horaria = st.number_input("Carga h/a", min_value=0, max_value=12, value=0, step=1)
        with col5:
            vagas = st.number_input("Vagas", min_value=1, max_value=20, value=1, step=1)

        descricao = st.text_input("Descrição/atividade", placeholder="Opcional. Deixe em branco para horário vago.")
        col6, col7 = st.columns(2)
        with col6:
            local = st.text_input("Local")
        with col7:
            habilitacao = st.text_input("Habilitação exigida")
        observacoes = st.text_area("Observações")
        salvar = st.form_submit_button("Adicionar horário vago")

    if salvar:
        if hora_fim <= hora_inicio:
            st.error("A hora final precisa ser maior que a hora inicial.")
        else:
            executar(
                """
                INSERT INTO horarios
                    (semana_id, data_aula, hora_inicio, hora_fim, disciplina, local, vagas, habilitacao, carga_horaria, observacoes)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    semana_id,
                    iso_date(data_aula),
                    hora_str(hora_inicio),
                    hora_str(hora_fim),
                    descricao.strip(),
                    local.strip(),
                    int(vagas),
                    habilitacao.strip(),
                    int(carga_horaria),
                    observacoes.strip(),
                ),
            )
            st.success("Horário vago adicionado.")
            st.rerun()

    st.divider()
    quadro = quadro_semana(semana_id)
    if quadro.empty:
        st.info("Nenhum horário cadastrado para esta semana.")
        return

    quadro_view = preparar_df_quadro(quadro)
    st.dataframe(quadro_view, use_container_width=True, hide_index=True)

    st.subheader("Editar, bloquear, liberar ou excluir horário")
    opcoes = {
        f"{row['horario_id']} | {br_date(row['data_aula'])} {periodo_com_carga(row['hora_inicio'], row['hora_fim'], row.get('carga_horaria', 0))} - {descricao_aula(row['disciplina'])}": int(row["horario_id"])
        for _, row in quadro.iterrows()
    }
    escolha = st.selectbox("Horário", list(opcoes.keys()))
    horario_id = opcoes[escolha]
    horario = consultar_um("SELECT * FROM horarios WHERE id = ?", (horario_id,))
    if horario:
        with st.form("editar_horario"):
            col1, col2, col3 = st.columns(3)
            with col1:
                novo_bloqueado = st.checkbox("Bloqueado", value=bool(horario["bloqueado"]))
            with col2:
                novas_vagas = st.number_input("Vagas", min_value=1, max_value=20, value=int(horario["vagas"]), step=1)
            with col3:
                nova_carga = st.number_input("Carga h/a", min_value=0, max_value=12, value=int(horario["carga_horaria"] or 0), step=1)
            nova_descricao = st.text_input("Descrição/atividade", value=horario["disciplina"] or "")
            col4, col5 = st.columns(2)
            with col4:
                novo_local = st.text_input("Local", value=horario["local"] or "")
            with col5:
                nova_habilitacao = st.text_input("Habilitação exigida", value=horario["habilitacao"] or "")
            novas_obs = st.text_area("Observações", value=horario["observacoes"] or "")
            salvar_h = st.form_submit_button("Salvar alterações do horário")

        if salvar_h:
            executar(
                """
                UPDATE horarios
                SET bloqueado = ?, vagas = ?, carga_horaria = ?, disciplina = ?, local = ?, habilitacao = ?, observacoes = ?
                WHERE id = ?
                """,
                (
                    1 if novo_bloqueado else 0,
                    int(novas_vagas),
                    int(nova_carga),
                    nova_descricao.strip(),
                    novo_local.strip(),
                    nova_habilitacao.strip(),
                    novas_obs.strip(),
                    horario_id,
                ),
            )
            st.success("Horário atualizado.")
            st.rerun()
        if st.button("Excluir horário", type="secondary"):
            executar("DELETE FROM horarios WHERE id = ?", (horario_id,))
            st.warning("Horário excluído.")
            st.rerun()


def aba_escolhas() -> None:
    st.subheader("Escolhas registradas")
    semana_id = selecionar_semana("Semana")
    if not semana_id:
        return

    df = consultar(
        """
        SELECT
            e.id AS escolha_id,
            h.id AS horario_id,
            h.data_aula,
            h.hora_inicio,
            h.hora_fim,
            h.disciplina,
            h.local,
            h.carga_horaria,
            TRIM(COALESCE(i.posto_grad, '') || ' ' || i.nome) AS instrutor,
            i.matricula,
            e.criado_em
        FROM escolhas e
        JOIN horarios h ON h.id = e.horario_id
        JOIN instrutores i ON i.id = e.instrutor_id
        WHERE h.semana_id = ?
          AND e.status = 'Confirmada'
        ORDER BY h.data_aula, h.hora_inicio, instrutor
        """,
        (semana_id,),
    )

    if df.empty:
        st.info("Nenhuma escolha registrada nesta semana.")
    else:
        df_view = df.assign(
            Data=lambda d: pd.to_datetime(d["data_aula"]).dt.strftime("%d/%m/%Y"),
            Horário=lambda d: d.apply(lambda r: periodo_com_carga(r["hora_inicio"], r["hora_fim"], r.get("carga_horaria", 0)), axis=1),
            Descrição=lambda d: d["disciplina"].apply(descricao_aula),
        )

        st.dataframe(
            df_view[["escolha_id", "Data", "Horário", "Descrição", "local", "instrutor", "matricula", "criado_em"]].rename(
                columns={
                    "escolha_id": "ID",
                    "local": "Local",
                    "instrutor": "Instrutor",
                    "matricula": "Matrícula",
                    "criado_em": "Registrado em",
                }
            ),
            use_container_width=True,
            hide_index=True,
        )

    st.divider()
    st.subheader("Adicionar escolha manualmente")
    st.caption("Use esta opção para o administrador ajustar o quadro sem depender do link público. Ainda são validados vagas, conflitos e limite de carga.")

    quadro = quadro_semana(semana_id)
    instrutores = listar_instrutores(ativos=True)
    if quadro.empty:
        st.warning("Cadastre horários antes de adicionar escolhas manualmente.")
    elif instrutores.empty:
        st.warning("Cadastre instrutores ativos antes de adicionar escolhas manualmente.")
    else:
        opcoes_h = {
            f"{row['horario_id']} | {br_date(row['data_aula'])} {periodo_com_carga(row['hora_inicio'], row['hora_fim'], row.get('carga_horaria', 0))} - {descricao_aula(row['disciplina'])} - vagas {int(row['ocupadas'])}/{int(row['vagas'])}": int(row["horario_id"])
            for _, row in quadro.iterrows()
        }
        opcoes_i = {f"{row['id']} | {formatar_instrutor(row)} - {row['matricula'] or 'sem matrícula'}": int(row["id"]) for _, row in instrutores.iterrows()}
        with st.form("adicionar_escolha_admin"):
            horario_label = st.selectbox("Horário", list(opcoes_h.keys()))
            instrutor_label = st.selectbox("Instrutor", list(opcoes_i.keys()))
            adicionar = st.form_submit_button("Adicionar escolha")
        if adicionar:
            ok, msg = escolher_horario(opcoes_i[instrutor_label], opcoes_h[horario_label], origem_admin=True)
            if ok:
                st.success(msg)
            else:
                st.error(msg)
            st.rerun()

    if not df.empty:
        st.divider()
        st.subheader("Remover escolha manualmente")
        opcoes = {
            f"{row['escolha_id']} | {br_date(row['data_aula'])} {periodo_com_carga(row['hora_inicio'], row['hora_fim'], row.get('carga_horaria', 0))} - {row['instrutor']} - {descricao_aula(row['disciplina'])}": int(row["escolha_id"])
            for _, row in df.iterrows()
        }
        escolha = st.selectbox("Escolha", list(opcoes.keys()))
        if st.button("Remover escolha selecionada"):
            executar("DELETE FROM escolhas WHERE id = ?", (opcoes[escolha],))
            st.success("Escolha removida.")
            st.rerun()

def aba_quadro_final() -> None:
    st.subheader("Quadro semanal final")
    semana_id = selecionar_semana("Semana do quadro")
    if not semana_id:
        return
    semana = obter_semana(semana_id)
    assert semana is not None

    quadro = quadro_semana(semana_id)
    if quadro.empty:
        st.info("Nenhum horário cadastrado para esta semana.")
        return

    total_horarios = len(quadro)
    vagas_total = int(quadro["vagas"].sum())
    vagas_ocupadas = int(quadro["ocupadas"].sum())
    pendentes = vagas_total - vagas_ocupadas
    carga_total = int((quadro["carga_horaria"].fillna(0) * quadro["vagas"].fillna(0)).sum())

    col1, col2, col3, col4, col5 = st.columns(5)
    col1.metric("Horários", total_horarios)
    col2.metric("Vagas totais", vagas_total)
    col3.metric("Vagas ocupadas", vagas_ocupadas)
    col4.metric("Pendências", pendentes)
    col5.metric("Carga total", f"{carga_total}h/a")
    st.caption(f"Status da semana: {semana['status']} | Limite por instrutor: {'sem limite' if limite_horas_da_semana(semana) <= 0 else str(limite_horas_da_semana(semana)) + 'h/a'}")

    quadro_view = preparar_df_quadro(quadro)
    st.dataframe(quadro_view, use_container_width=True, hide_index=True)

    cargas = carga_por_instrutor_semana(semana_id)
    if not cargas.empty:
        st.markdown("### Carga por instrutor")
        st.dataframe(
            cargas.rename(columns={"instrutor": "Instrutor", "matricula": "Matrícula", "quantidade_horarios": "Horários", "carga_total": "Carga h/a"})[["Instrutor", "Matrícula", "Horários", "Carga h/a"]],
            use_container_width=True,
            hide_index=True,
        )

    col_fechar, col_publicar = st.columns(2)
    with col_fechar:
        if st.button("Encerrar semana a partir deste quadro"):
            executar("UPDATE semanas SET status = 'Encerrada', encerrada_em = ? WHERE id = ?", (datetime.now().isoformat(timespec="minutes"), semana_id))
            st.success("Semana encerrada.")
            st.rerun()
    with col_publicar:
        if st.button("Publicar semana a partir deste quadro"):
            executar("UPDATE semanas SET status = 'Publicada', publicada_em = ? WHERE id = ?", (datetime.now().isoformat(timespec="minutes"), semana_id))
            st.success("Semana publicada.")
            st.rerun()

    titulo = f"{semana['titulo']} - {br_date(semana['data_inicio'])} a {br_date(semana['data_fim'])}"
    excel_bytes = exportar_excel(quadro_view, titulo)
    st.download_button(
        "Baixar quadro em Excel",
        data=excel_bytes,
        file_name=f"quadro_instrutores_semana_{semana_id}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    if REPORTLAB_DISPONIVEL:
        pdf_bytes = exportar_pdf(quadro_view, titulo)
        st.download_button(
            "Baixar quadro em PDF",
            data=pdf_bytes,
            file_name=f"quadro_instrutores_semana_{semana_id}.pdf",
            mime="application/pdf",
        )
    else:
        st.warning("PDF indisponível: instale a biblioteca reportlab com `pip install reportlab`.")

    texto = gerar_texto_whatsapp(quadro, semana)
    st.text_area("Texto para WhatsApp", value=texto, height=280)


def aba_link_publico() -> None:
    st.subheader("Link público para escolha dos instrutores")
    semana_id = selecionar_semana("Semana")
    if not semana_id:
        return
    semana = obter_semana(semana_id)
    assert semana is not None

    base_url = st.text_input(
        "Endereço público do app",
        value=obter_base_url(),
        help="Quando publicar o app, troque pelo endereço público. Localmente, use http://localhost:8501.",
    ).rstrip("/")
    link = f"{base_url}/?pagina=escolha&token={semana['token']}"

    st.code(link, language="text")
    st.caption("Envie este link aos instrutores. Eles não precisam acessar a área administrativa.")

    st.divider()
    st.write("Mensagem pronta para enviar:")
    mensagem = (
        f"Senhores instrutores, segue o link para escolha dos horários da semana:\n\n"
        f"{semana['titulo']}\n"
        f"Período: {br_date(semana['data_inicio'])} a {br_date(semana['data_fim'])}\n"
    )
    if semana["prazo_escolha"]:
        mensagem += f"Prazo: {datetime.fromisoformat(semana['prazo_escolha']).strftime('%d/%m/%Y %H:%M')}\n"
    mensagem += f"\n{link}"
    st.text_area("Mensagem", value=mensagem, height=180)


# =============================================================================
# Roteamento
# =============================================================================


def main() -> None:
    init_db()
    pagina = get_query_param("pagina", "admin")
    token = get_query_param("token", "")

    if pagina == "escolha" and token:
        pagina_escolha_publica(token)
    else:
        pagina_admin()


if __name__ == "__main__":
    main()
